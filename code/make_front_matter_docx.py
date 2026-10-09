# -*- coding: utf-8 -*-
"""生成 Cell Reports Methods 投稿配套文件：Highlights + eTOC Blurb（单个 Word 文档）。

要求（Final file requirements）：
- Highlights：3–4 条，每条 <=85 字符（含空格）
- eTOC Blurb：<=50 词，第三人称，以 "First Author et al." 开头
两个部分放在同一个 Word 文档中上传。

用法：python code/make_front_matter_docx.py
"""
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt

OUT_DIR = Path(__file__).resolve().parent.parent / 'paper' / 'CellReportsMethods'

HIGHLIGHTS = [
    'Random trial subsampling makes EEG decoding statistics seed-dependent',
    'Single-seed permutation tests inflate Type I error under stochastic decoding',
    'A seed-aggregated statistic restores calibrated permutation inference',
    'An audit protocol separates stability from significance in EEG decoding',
]

ETOC = (
    'Tan et al. show that randomized trial subsampling makes single-trial EEG '
    'decoding statistics seed-dependent, so single-seed permutation tests can '
    'be unreliable. They introduce a seed-aggregated statistic with calibrated '
    'permutation inference and an audit protocol reporting stability and '
    'significance separately, retracting a previously flagged decoding effect '
    'in food-EEG data.'
)


def main():
    # --- 自检：硬性字数约束 ---
    for i, h in enumerate(HIGHLIGHTS, 1):
        n = len(h)
        flag = 'OK' if n <= 85 else 'TOO LONG'
        print(f'  highlight {i}: {n:3d} chars [{flag}]  {h}')
        if n > 85:
            sys.exit(f'ERROR: highlight {i} 超过 85 字符')
    n_words = len(ETOC.split())
    print(f'  eTOC: {n_words} words ({"OK" if n_words <= 50 else "TOO LONG"})')
    if n_words > 50:
        sys.exit('ERROR: eTOC blurb 超过 50 词')
    if not (3 <= len(HIGHLIGHTS) <= 4):
        sys.exit('ERROR: Highlights 必须为 3–4 条')

    doc = Document()
    doc.add_heading('Highlights', level=1)
    for h in HIGHLIGHTS:
        p = doc.add_paragraph(h, style='List Bullet')
        p.runs[0].font.size = Pt(11)

    doc.add_heading('eTOC Blurb', level=1)
    p = doc.add_paragraph(ETOC)
    p.runs[0].font.size = Pt(11)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out = OUT_DIR / 'Highlights_and_eTOC.docx'
    doc.save(str(out))
    print(f'[saved] {out}')


if __name__ == '__main__':
    main()
