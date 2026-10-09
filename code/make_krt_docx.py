# -*- coding: utf-8 -*-
"""生成 Cell Reports Methods 投稿配套文件：Key Resources Table (KRT)。

Cell Press KRT 三列：REAGENT or RESOURCE | SOURCE | IDENTIFIER。
凡本研究中未使用/未生成的类别，显式写明 "none"；IDENTIFIER 只填写可核验的
版本号 / 访问号 / 网址，不臆造 RRID。

用法：python code/make_krt_docx.py
"""
from pathlib import Path

from docx import Document
from docx.shared import Pt

OUT_DIR = Path(__file__).resolve().parent.parent / 'paper' / 'CellReportsMethods'

GH = 'https://github.com/fatbaby614/TasteHealthToFood'

# (类别标题, [(资源, 来源, 标识符), ...])
SECTIONS = [
    ('Antibodies', [('none', 'n/a', 'n/a')]),
    ('Bacterial and virus strains', [('none', 'n/a', 'n/a')]),
    ('Biological samples', [('none', 'n/a', 'n/a')]),
    ('Chemicals, peptides, and recombinant proteins', [('none', 'n/a', 'n/a')]),
    ('Critical commercial assays', [('none', 'n/a', 'n/a')]),
    ('Deposited data', [
        ('FoodEEG dataset (64-channel EEG, 117 participants; food '
         'categorisation and food-choice tasks)',
         'Chae et al., 2025; OpenNeuro',
         'OpenNeuro: ds007012, version 1.1.0; '
         'https://openneuro.org/datasets/ds007012/versions/1.1.0'),
        ('Per-seed decoding statistics, permutation distributions, '
         'calibration outputs, and sensitivity tables',
         'This paper', f'Archived in the code repository: {GH}'),
    ]),
    ('Experimental models: Cell lines', [('none', 'n/a', 'n/a')]),
    ('Experimental models: Organisms/strains', [('none', 'n/a', 'n/a')]),
    ('Oligonucleotides', [('none', 'n/a', 'n/a')]),
    ('Recombinant DNA', [('none', 'n/a', 'n/a')]),
    ('Software and algorithms', [
        ('Python', 'Python Software Foundation',
         '3.13.9; https://www.python.org/'),
        ('NumPy', 'Harris et al., 2020', '2.3.5; https://numpy.org/'),
        ('SciPy', 'Virtanen et al., 2020', '1.17.1; https://scipy.org/'),
        ('scikit-learn', 'Pedregosa et al., 2011',
         '1.7.2; https://scikit-learn.org/'),
        ('joblib', 'Joblib developers', '1.5.2; https://joblib.readthedocs.io/'),
        ('pyriemann', 'Barachant et al.', '0.11; https://pyriemann.readthedocs.io/'),
        ('pandas', 'Pandas developers', '2.3.3; https://pandas.pydata.org/'),
        ('MNE-Python', 'Gramfort et al., 2013', '1.12.1; https://mne.tools/'),
        ('Matplotlib', 'Hunter, 2007', '3.10.6; https://matplotlib.org/'),
        ('HSSM', 'Fengler et al., 2026', '0.3.0; https://github.com/lnccbrown/HSSM'),
        ('PyMC', 'Abril-Pla et al., 2023', '6.0.1; https://www.pymc.io/'),
        ('ArviZ', 'ArviZ developers', '1.2.0; https://python.arviz.org/'),
        ('Seed-audit, calibration, and sensitivity analysis code; '
         'frozen analysis configuration; figure-generation scripts',
         'This paper', f'{GH}'),
    ]),
    ('Other', [('none', 'n/a', 'n/a')]),
]


def main():
    doc = Document()
    doc.add_heading('Key Resources Table', level=1)

    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    for c, txt in zip(hdr, ('REAGENT or RESOURCE', 'SOURCE', 'IDENTIFIER')):
        c.text = ''
        r = c.paragraphs[0].add_run(txt)
        r.bold = True
        r.font.size = Pt(9)

    n_rows = 0
    for title, items in SECTIONS:
        row = table.add_row().cells
        row[0].text = ''
        r = row[0].paragraphs[0].add_run(title)
        r.italic = True
        r.bold = True
        r.font.size = Pt(9)
        for res, src, ident in items:
            row = table.add_row().cells
            for cell, txt in zip(row, (res, src, ident)):
                cell.text = txt
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(9)
            n_rows += 1

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out = OUT_DIR / 'Key_Resources_Table.docx'
    doc.save(str(out))
    print(f'[saved] {out}  ({len(SECTIONS)} 类别, {n_rows} 资源行)')


if __name__ == '__main__':
    main()
